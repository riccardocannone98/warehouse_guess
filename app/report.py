"""Vista a matrice (prodotto × taglia) ed export Excel dei risultati."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .sku import size_key


def build_matrix(rows) -> list[dict]:
    """Raggruppa i risultati per codice; per ogni codice una riga per prodotto/sito
    e una colonna per taglia."""
    by_code: dict[str, dict] = {}
    for r in rows:
        block = by_code.setdefault(r["code"], {"code": r["code"], "sizes": set(), "products": {}})
        block["sizes"].add(r["size"])
        key = (r["site_name"], r["product_url"])
        product = block["products"].setdefault(
            key,
            {
                "site_name": r["site_name"],
                "title": r["product_title"],
                "url": r["product_url"],
                "image": r["image"],
                "currency": r["currency"],
                "price": r["price"],
                "cells": {},
                "total": 0,
                "unknown": 0,
            },
        )
        if r["price"] is not None and (product["price"] is None or r["price"] < product["price"]):
            product["price"] = r["price"]
        cell = product["cells"].setdefault(r["size"], {"qty": None, "available": None, "method": r["method"]})
        if r["quantity"] is not None:
            cell["qty"] = (cell["qty"] or 0) + r["quantity"]
            product["total"] += r["quantity"]
        if r["available"] is not None:
            cell["available"] = bool(r["available"]) or bool(cell["available"])
        if r["quantity"] is None and r["available"] != 0:
            product["unknown"] += 1

    blocks = []
    for block in by_code.values():
        sizes = sorted(block["sizes"], key=size_key)
        products = sorted(block["products"].values(), key=lambda p: (-p["total"], p["site_name"]))
        size_totals = {s: sum((p["cells"].get(s) or {}).get("qty") or 0 for p in products) for s in sizes}
        blocks.append({"code": block["code"], "sizes": sizes, "products": products, "size_totals": size_totals})
    return blocks


HEADERS = [
    ("Codice prodotto", 18),
    ("Sito", 22),
    ("Prodotto", 42),
    ("Taglia", 10),
    ("Quantità [pz]", 14),
    ("Disponibile [sì/no]", 18),
    ("Metodo rilevazione", 22),
    ("Prezzo [valuta sito]", 18),
    ("Valuta", 8),
    ("SKU variante", 22),
    ("URL prodotto", 60),
]


def export_xlsx(search, log, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dati"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")

    ws.append([h for h, _ in HEADERS])
    for i, (_, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = header_font, header_fill
        ws.column_dimensions[cell.column_letter].width = width
    for r in sorted(rows, key=lambda r: (r["code"], r["site_name"], r["product_title"], size_key(r["size"]))):
        available = None if r["available"] is None else ("sì" if r["available"] else "no")
        ws.append([
            r["code"], r["site_name"], r["product_title"], r["size"], r["quantity"], available,
            r["method"], r["price"], r["currency"], r["variant_sku"], r["product_url"],
        ])
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    log_ws = wb.create_sheet("Log ricerca")
    log_ws.append(["Codice prodotto", "Sito", "Esito", "Dettaglio"])
    for i, width in enumerate((18, 22, 10, 80), start=1):
        cell = log_ws.cell(row=1, column=i)
        cell.font, cell.fill = header_font, header_fill
        log_ws.column_dimensions[cell.column_letter].width = width
    for entry in log:
        log_ws.append([entry["code"], entry["site_name"], entry["status"], entry["message"]])
    log_ws.append([])
    log_ws.append(["Data ricerca", search["created_at"]])
    log_ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
