import re
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from app.sku import parse_codes, qty_from_text, size_key

from .fakes import BLOCKED, GENERIC, SHOP, WOO, FakeInternet


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path / "test.db", FakeInternet().client_factory)
    return TestClient(app)


def test_parse_codes_dedupes_and_normalizes():
    assert parse_codes("dd1391-100\nDD1391 100, ie3438; \n") == ["DD1391-100", "IE3438"]


def test_size_sorting():
    assert sorted(["44", "40 ½", "42.5", "XL", "40"], key=size_key) == ["40", "40 ½", "42.5", "44", "XL"]


def test_qty_from_text():
    assert qty_from_text("Solo 3 rimasti!") == 3
    assert qty_from_text("2 disponibili") == 2
    assert qty_from_text("Only 1 left") == 1
    assert qty_from_text("Disponibile") is None


def test_sites_crud_and_import(client):
    r = client.post("/sites", data={"base_url": "shop.example/", "platform": "auto", "active": "on"})
    assert r.status_code == 200 and "Aggiunto shop.example" in r.text
    r = client.post("/sites/import", data={"text": f"Woo;{WOO}\n{SHOP}\nhttps://nuovo.example"})
    assert "Importati 2 siti (1 già presenti)" in r.text
    page = client.get("/sites").text
    assert "https://nuovo.example" in page and "Woo" in page
    csv_text = client.get("/sites/export.csv").content.decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("nome;url;piattaforma")
    assert len(csv_text.strip().splitlines()) == 4


def test_invalid_url_rejected(client):
    assert client.post("/sites", data={"base_url": "nonunurl", "platform": "auto"}).status_code == 400


def test_full_search_flow(client):
    client.post("/sites/import", data={"text": f"Shop;{SHOP}\nWoo;{WOO}\nAtelier;{GENERIC}\nBloccato;{BLOCKED}"})
    r = client.post("/search", data={"codes": "DD1391-100\nB75806\nU9060AAA"})
    assert r.status_code == 200 and "/results/" in str(r.url)
    html = r.text
    assert "Nike Dunk Low Retro White Black" in html
    assert "adidas Samba OG Cloud White" in html
    assert "New Balance 9060 Sea Salt" in html
    assert "HTTP 403" in html  # sito bloccato riportato nel log

    # La piattaforma rilevata viene salvata sul sito.
    sites = client.get("/sites").text
    assert "Shopify" in sites and "WooCommerce" in sites and "Generico" in sites

    search_id = str(r.url).rsplit("/", 1)[1]
    xlsx = client.get(f"/results/{search_id}/export.xlsx")
    wb = load_workbook(BytesIO(xlsx.content))
    rows = list(wb["Dati"].iter_rows(values_only=True))
    assert rows[0][:5] == ("Codice prodotto", "Sito", "Prodotto", "Taglia", "Quantità [pz]")
    samba = {r[3]: r[4] for r in rows[1:] if r[0] == "B75806"}
    assert samba == {"41.5": 1, "43": 7, "45": 0}
    assert any(r[2] == "errore" for r in wb["Log ricerca"].iter_rows(values_only=True))

    dash = client.get("/").text
    assert "DD1391-100 · B75806 · U9060AAA" in dash  # ricerche recenti
    assert "Dunk Low Retro" in dash  # nome modello dal catalogo
    assert "Taglia 43: 7 pz" in dash  # grafico pezzi per taglia


def test_search_without_sites_or_codes(client):
    assert client.post("/search", data={"codes": "DD1391-100"}).status_code == 400
    client.post("/sites", data={"base_url": SHOP, "platform": "shopify", "active": "on"})
    assert client.post("/search", data={"codes": "  "}).status_code == 400


def test_search_only_selected_sites(client):
    client.post("/sites/import", data={"text": f"Shop;{SHOP}\nWoo;{WOO}"})
    r = client.post("/search", data={"codes": "DD1391-100", "site_ids": ["1"]})
    assert "Shop" in r.text and "Woo</td>" not in r.text


def test_dashboard_empty_state(client):
    page = client.get("/").text
    assert "Ancora nessuna rilevazione" in page


def test_dashboard_trend_after_two_searches(client):
    client.post("/sites", data={"base_url": WOO, "platform": "woocommerce", "active": "on"})
    client.post("/search", data={"codes": "B75806"})
    client.post("/search", data={"codes": "b75806"})  # stesso codice scritto diversamente
    page = client.get("/").text
    assert page.count("<polyline") == 1 and "+0" in page


def test_catalog_seeded_and_crud(client):

    page = client.get("/catalog").text
    assert "DD1391-100" in page and "da verificare" in page

    r = client.post("/catalog", data={"brand": "Asics", "model": "Gel-Kayano 14", "colorway": "", "code": "1201a019 107"})
    assert "Aggiunto 1201A019 107" in r.text
    r = client.post("/catalog", data={"brand": "Nike", "model": "Dunk", "colorway": "", "code": "dd1391100"})
    assert "già in catalogo" in r.text

    r = client.post("/catalog/import", data={"text": "Nike;Air Max 90;White;CN8490-100\nadidas;Samba OG;B75806\nincompleta"})
    assert "Importati 1 codici (2 già presenti o incompleti)" in r.text

    item_id = re.search(r"/catalog/(\d+)/delete\" onsubmit=\"return confirm\('Eliminare 1201A019 107", r.text).group(1)
    r = client.post(f"/catalog/{item_id}", data={"brand": "ASICS", "model": "Gel-Kayano 14", "colorway": "Cream", "code": "1201A019-107"})
    assert "Aggiornato 1201A019-107" in r.text and "Cream" in r.text
    client.post(f"/catalog/{item_id}/delete")
    assert "1201A019-107" not in client.get("/catalog").text

    csv_text = client.get("/catalog/export.csv").content.decode("utf-8-sig")
    assert csv_text.startswith("marca;modello;colorway;codice") and "CN8490-100" in csv_text


def test_search_page_shows_catalog_and_prefills(client):
    page = client.get("/search", params=[("codes", "dd1391-100"), ("codes", "B75806")]).text
    assert 'data-code="CW2288-111"' in page
    assert 'value="DD1391-100\nB75806"' in page


def test_search_with_codes_selected_from_catalog(client):
    client.post("/sites", data={"base_url": SHOP, "platform": "shopify", "active": "on"})
    r = client.post("/search", data={"codes": ["DD1391-100", "CW2288-111"]})
    assert "/results/" in str(r.url)
    assert "Nike Dunk Low Retro White Black" in r.text and "CW2288-111" in r.text
